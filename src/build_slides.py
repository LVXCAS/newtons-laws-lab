"""Build lab slides (graph + condensed table per trial, plus a summary per law) as .pptx + condensed CSVs."""
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

sys.path.insert(0, os.path.dirname(__file__))
import build_angled as ang  # noqa: E402
import build_n3l as n3l  # noqa: E402
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from build_n1l import coast_stats, write_csv  # noqa: E402
from build_n2l import CART_KG, DOCS, TRIALS, WEIGHT_KG, analyze, linfit  # noqa: E402
from export_gambl import read_runs  # noqa: E402

OUT_DIR = os.environ.get("OUT_DIR", os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "outputs"))
N1L_G = os.path.join(OUT_DIR, "n1l", "Graphs")
N2L_G = os.path.join(OUT_DIR, "n2l", "Graphs")
N3L_G = os.path.join(OUT_DIR, "n3l", "Graphs")
ANG_G = os.path.join(OUT_DIR, "angled", "Graphs")
HEADER_RGB = RGBColor(0x1F, 0x4E, 0x79)
EVENT_RGB = RGBColor(0xFD, 0xE9, 0xD9)


def nearest(pts, t):
    return min(range(len(pts)), key=lambda i: abs(pts[i][0] - t))


def whole_seconds(pts):
    """Indices of the samples closest to 0 s, 1 s, 2 s ... across the recording."""
    return [nearest(pts, s) for s in range(0, int(pts[-1][0]) + 1)]


def condensed(pts, events, fmt):
    """Whole-second rows plus labelled event rows, in time order, without duplicates."""
    rows = {i: "" for i in whole_seconds(pts)}
    for label, i in events:
        rows[i] = label
    return [fmt(i, rows[i]) for i in sorted(rows)]


def f(x, n=3):
    return f"{x:.{n}f}"


SINGLE_HDR = ["Time (s)", "Position (m)", "Velocity (m/s)", "Accel. (m/s²)", "Phase"]


def single_car_fmt(pts):
    def fmt(i, label):
        p = pts[i]
        return [f(p[0], 2), f(p[1]), f(p[2]), f(p[3]), label]
    return fmt


def n1l_n2l_data():
    runs = read_runs(os.path.join(DOCS, "N2L_5.gambl"))
    rows = [{"name": name, "weights": w, **analyze(runs[num - 1])} for name, num, w in TRIALS]
    return rows, [coast_stats(row) for row in rows]


def n1l_slides(rows, stats):
    slides = []
    for k, (row, cs) in enumerate(zip(rows, stats), start=1):
        pts = row["pts"]
        c0, c1 = min(row["coast"]), max(row["coast"])
        mid = nearest(pts, (pts[c0][0] + pts[c1][0]) / 2)
        events = [("Coasting starts", c0), ("Coasting (middle)", mid), ("Late coasting", max(c1 - 2, mid))]
        mass = CART_KG + row["weights"] * WEIGHT_KG
        slides.append({
            "title": f"N1L Trial {k}",
            "subtitle": f"Car 1, mass {mass:.3f} kg ({row['weights']} weights) — coasting at nearly constant velocity",
            "image": os.path.join(N1L_G, f"N1L_{k} - coasting position and velocity.png"),
            "header": SINGLE_HDR, "rows": condensed(pts, events, single_car_fmt(pts)),
            "note": f"Avg coasting v = {cs['v_avg']:.3f} m/s; coasting a = {cs['a']:.3f} m/s² (≈ 0, small friction)"})
    avg_a = sum(cs["a"] for cs in stats) / len(stats)
    summary = {
        "title": "N1L Summary",
        "subtitle": "Once the push stops, each car keeps a nearly constant velocity (net force ≈ 0)",
        "image": os.path.join(N1L_G, "N1L Results - coasting velocity all trials.png"),
        "header": ["Trial", "Mass (kg)", "Avg v (m/s)", "% change in v", "Coast a (m/s²)", "R² of x-t"],
        "rows": [[f"Trial {k}", f(CART_KG + row["weights"] * WEIGHT_KG), f(cs["v_avg"]), f(cs["pct"], 1),
                  f(cs["a"]), f(cs["x_r2"], 4)] for k, (row, cs) in enumerate(zip(rows, stats), start=1)],
        "note": f"Average coasting a = {avg_a:.3f} m/s² — tiny compared with the launch (3–5 m/s²)"}
    return slides + [summary]


def n2l_slides(rows):
    slides = []
    for k, row in enumerate(rows, start=1):
        pts = row["pts"]
        p0, p1 = min(row["push"]), max(row["push"])
        events = [("Launch starts", p0), ("Launch ends", p1), ("Coasting", min(row["coast"]))]
        mass = CART_KG + row["weights"] * WEIGHT_KG
        slides.append({
            "title": f"N2L Trial {k}",
            "subtitle": f"Car 1, mass {mass:.3f} kg ({row['weights']} weights) — same plunger launch",
            "image": os.path.join(N2L_G, f"N2L_{k} - velocity and acceleration.png"),
            "header": SINGLE_HDR, "rows": condensed(pts, events, single_car_fmt(pts)),
            "note": f"Launch acceleration (slope of v-t during launch) = {row['a_push']:.3f} m/s²"})
    masses = [CART_KG + row["weights"] * WEIGHT_KG for row in rows]
    accs = [row["a_push"] for row in rows]
    slope, icpt, r2 = linfit([1 / m for m in masses], accs)
    summary = {
        "title": "N2L Summary",
        "subtitle": "As mass increases, acceleration decreases — a vs 1/m is a straight line",
        "image": os.path.join(N2L_G, "N2L Results - a vs m and a vs 1-m.png"),
        "header": ["Trial", "Mass m (kg)", "1/m (1/kg)", "Launch a (m/s²)", "m·a (N)"],
        "rows": [[f"Trial {k}", f(m), f(1 / m), f(a), f(m * a)]
                 for k, (m, a) in enumerate(zip(masses, accs), start=1)],
        "note": f"Best fit: a = {slope:.3f}(1/m) + {icpt:.3f},  R² = {r2:.3f}"}
    return slides + [summary]


def n3l_fmt(res):
    pts = res["pts"]

    def fmt(i, label):
        p = pts[i]
        return [f(p[0], 2), f(p[2]), f(res["flip"] * p[1]), f(res["fg"][i]), f(res["fy"][i]), label]
    return fmt


def n3l_slides():
    runs = read_runs(os.path.join(DOCS, "N3L_6.gambl"))
    slides, results = [], []
    for k, (num, src, ny, ng, kind) in enumerate(n3l.RUNS, start=1):
        res = n3l.analyze(runs[num - 1], ny, ng)
        results.append((k, num, res))
        lo, hi = res["lo"], res["hi"]
        peak = max(range(lo, hi + 1), key=lambda i: abs(res["fg"][i]))
        events = [("Before contact", max(lo - 1, 0)), ("Contact (peak force)", peak),
                  ("After contact", min(hi + 1, len(res["pts"]) - 1))]
        slides.append({
            "title": f"N3L Trial {k}",
            "subtitle": f"Run {num} ({src}) — {kind};  m₁ = {res['mg']:.3f} kg, m₂ = {res['my']:.3f} kg",
            "image": os.path.join(N3L_G, f"N3L Run {num:02d} ({src}) - velocity and force.png"),
            "header": ["Time (s)", "v₁ (m/s)", "v₂ (m/s)", "F on Car 1 (N)", "F on Car 2 (N)", "Phase"],
            "rows": condensed(res["pts"], events, n3l_fmt(res)),
            "note": f"Impulse on Car 1 = {res['jg']:.4f} N·s, on Car 2 = {res['jy']:.4f} N·s "
                    f"({n3l.pct_diff(res['jg'], res['jy']):.1f}% apart, opposite signs)"})
    summary = {
        "title": "N3L Summary",
        "subtitle": "In every collision the two cars push on each other equally and oppositely",
        "image": os.path.join(N3L_G, "N3L Results - impulse on Y vs impulse on G.png"),
        "header": ["Trial", "m₁ (kg)", "m₂ (kg)", "Δv₁ (m/s)", "Δv₂ (m/s)", "J₁ (N·s)", "J₂ (N·s)", "% diff"],
        "rows": [[f"{k} (Run {num})", f(res["mg"]), f(res["my"]), f(res["dvg"]), f(res["dvy"]),
                  f(res["jg"], 4), f(res["jy"], 4), f(n3l.pct_diff(res["jg"], res["jy"]), 1)]
                 for k, num, res in results],
        "note": "J = m·Δv. Masses inferred from the weight pattern — confirm with lab notes."}
    return slides + [summary]


def angled_summary_graph(parts_by_run, a_up, a_down, theta, out_png):
    fig, ax = plt.subplots(figsize=(9, 5))
    labels, colors, vals = [], [], []
    for k, (num, parts) in enumerate(parts_by_run, start=1):
        for j, part in enumerate(parts, start=1):
            labels.append(f"T{k} roll {j}")
            colors.append("#C00000" if part["dir"] == "Up the ramp" else "#2E75B6")
            vals.append(part["a"])
    ax.bar(range(len(vals)), vals, color=colors)
    ax.axhline(a_up, color="#C00000", ls="--", lw=1.5, label=f"Avg rolling UP = {a_up:.3f} m/s²")
    ax.axhline(a_down, color="#2E75B6", ls="--", lw=1.5, label=f"Avg rolling DOWN = {a_down:.3f} m/s²")
    ax.set_xticks(range(len(vals)), labels, rotation=30, fontsize=8)
    ax.set_ylabel("Acceleration (m/s²)  (+ = up the ramp)")
    ax.set_ylim(min(vals) * 1.25, 0.05)
    ax.set_title(f"Angled ramp — same constant acceleration every roll (ramp ≈ {theta:.1f}°)")
    ax.legend(loc="lower right")
    ax.grid(axis="y", alpha=0.3)
    fig.tight_layout()
    fig.savefig(out_png, dpi=120)
    plt.close(fig)


def angled_slides():
    import math
    runs = read_runs(os.path.join(DOCS, "ANGLED.gambl"))
    slides, parts_by_run, ups, downs = [], [], [], []
    for k, (num, kind) in enumerate(ang.RUNS, start=1):
        pts, parts = ang.analyze(runs[num - 1])
        parts_by_run.append((num, parts))
        events = [(f"Roll {j} starts ({p['dir'].split()[0].lower()})", p["idx"][0]) for j, p in enumerate(parts, start=1)]
        for p in parts:
            (ups if p["dir"] == "Up the ramp" else downs).append(p["a"])
        slides.append({
            "title": f"Angled Trial {k}",
            "subtitle": f"Run {num} (ANGLED.gambl), Car 1 — {kind.lower()}",
            "image": os.path.join(ANG_G, f"ANGLED Run {num} - position and velocity.png"),
            "header": SINGLE_HDR, "rows": condensed(pts, events, single_car_fmt(pts)),
            "note": "Free-roll accelerations: " + ", ".join(f"{p['a']:.3f}" for p in parts) + " m/s²"})
    a_up, a_down = sum(ups) / len(ups), sum(downs) / len(downs)
    g_sin, friction = -(a_up + a_down) / 2, -(a_up - a_down) / 2
    theta = math.degrees(math.asin(g_sin / ang.G))
    summary_png = os.path.join(ANG_G, "ANGLED Results - acceleration every roll.png")
    angled_summary_graph(parts_by_run, a_up, a_down, theta, summary_png)
    rows = []
    for k, (num, parts) in enumerate(parts_by_run, start=1):
        for j, p in enumerate(parts, start=1):
            rows.append([f"Trial {k}", f"Roll {j}", p["dir"], f(p["t0"], 2) + "–" + f(p["t1"], 2), f(p["a"]), f(p["r2"], 4)])
    summary = {
        "title": "Angled Summary",
        "subtitle": "Gravity along the ramp gives a constant acceleration; friction makes 'up' slightly stronger than 'down'",
        "image": summary_png,
        "header": ["Trial", "Roll", "Direction", "Time (s)", "a (m/s²)", "R² of v-t"],
        "rows": rows,
        "note": f"g·sinθ = {g_sin:.3f} m/s² → ramp ≈ {theta:.1f}°;  friction ≈ {friction:.3f} m/s²;  mass cancels"}
    return slides + [summary]


def add_table(slide, header, rows, left, top, width):
    font = Pt(10) if len(rows) <= 10 else Pt(9)
    height = Inches(0.28) * (len(rows) + 1)
    table = slide.shapes.add_table(len(rows) + 1, len(header), left, top, width, height).table
    for c, text in enumerate(header):
        cell = table.cell(0, c)
        cell.text = text
        cell.fill.solid()
        cell.fill.fore_color.rgb = HEADER_RGB
        run = cell.text_frame.paragraphs[0].runs[0]
        run.font.size, run.font.bold = font, True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for r, row in enumerate(rows, start=1):
        highlight = header[-1] == "Phase" and row[-1]
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(text) if text != "" else " "
            cell.text_frame.paragraphs[0].runs[0].font.size = font
            if highlight:
                cell.fill.solid()
                cell.fill.fore_color.rgb = EVENT_RGB


def add_text(slide, text, left, top, width, height, size, bold=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    box.text_frame.word_wrap = True
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size, run.font.bold = Pt(size), bold


def build_pptx(all_slides, out_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    for s in all_slides:
        slide = prs.slides.add_slide(blank)
        add_text(slide, s["title"], Inches(0.4), Inches(0.2), Inches(12.5), Inches(0.7), 30, bold=True)
        add_text(slide, s["subtitle"], Inches(0.4), Inches(0.85), Inches(12.5), Inches(0.5), 14)
        slide.shapes.add_picture(s["image"], Inches(0.3), Inches(1.45), width=Inches(6.6))
        add_table(slide, s["header"], s["rows"], Inches(7.1), Inches(1.5), Inches(5.9))
        add_text(slide, s["note"], Inches(7.1), Inches(6.55), Inches(5.9), Inches(0.7), 12)
    prs.save(out_path)


def main(out_dir):
    os.makedirs(out_dir, exist_ok=True)
    rows, stats = n1l_n2l_data()
    groups = {"N1L": n1l_slides(rows, stats), "N2L": n2l_slides(rows), "N3L": n3l_slides(),
              "Angled": angled_slides()}
    for law, slides in groups.items():
        csv_dir = os.path.join(out_dir, f"{law} slide tables")
        os.makedirs(csv_dir, exist_ok=True)
        for s in slides:
            write_csv(os.path.join(csv_dir, f"{s['title']}.csv"), s["header"], s["rows"])
    all_slides = [s for law in groups.values() for s in law]
    for s in all_slides:
        assert os.path.exists(s["image"]), s["image"]
    build_pptx(all_slides, os.path.join(out_dir, "Newtons Laws Lab Slides.pptx"))
    print(len(all_slides), "slides")
    for s in all_slides[-3:]:
        print(s["title"], s["header"])
        for row in s["rows"]:
            print("   ", row)


if __name__ == "__main__":
    main(sys.argv[1])
